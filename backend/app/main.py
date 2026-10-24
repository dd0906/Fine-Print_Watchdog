import base64
import binascii
import hashlib
import hmac
import io
import logging
import os
import time
import uuid
import zipfile
from datetime import date
from pathlib import Path
from typing import Literal
from threading import Lock
from dotenv import load_dotenv
load_dotenv(Path(__file__).resolve().parents[2] / '.env')
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, model_validator
from .core import RULES, chunks, compare_text, deadline_candidates, document_deadline_candidates, cost_candidates, evidence, loan_cost, make_calendar, rule_findings, verify_candidates
from . import llm
from .privacy import DEFAULT_CATEGORIES as DEFAULT_PRIVACY_CATEGORIES, redact_pages, scan_pages
from .policies import CATALOGUE_VERSION, CATALOGUE_CHECKED_ON, POLICIES, catalogue_payload, match_policies

app = FastAPI(title='Fine-Print Watchdog', version='1.0.0')
logger = logging.getLogger(__name__)
DOCUMENTS = {}
STORE_LOCK = Lock()
MAX_BYTES = 8 * 1024 * 1024
MAX_CHARS = 120_000
MAX_DOCUMENTS_PER_SESSION = 20
MAX_DOCUMENTS_TOTAL = 200
DOCUMENT_TTL_SECONDS = 8 * 60 * 60


def cleanup_expired_documents():
    # Expire in-memory uploads after eight hours.
    cutoff = time.time() - DOCUMENT_TTL_SECONDS
    with STORE_LOCK:
        expired = [key for key, item in DOCUMENTS.items() if item.get('_created_at', 0) < cutoff]
        for key in expired:
            DOCUMENTS.pop(key, None)


@app.middleware('http')
async def optional_basic_auth(request: Request, call_next):
    """Protect hosted routes and assign each browser a private document namespace.

    Basic Auth is a deployment gate; the random HttpOnly session cookie separates
    documents between browsers even when teammates share the demo credentials.
    Render's health check remains public. Local mode stays easy to run, but still
    receives per-browser document isolation.
    """
    require_auth = os.getenv('REQUIRE_ACCESS_AUTH', 'false').strip().lower() in {'1', 'true', 'yes', 'on'}
    is_health = request.url.path == '/api/health'
    cleanup_expired_documents()

    if require_auth and not is_health:
        expected_user = os.getenv('WATCHDOG_AUTH_USERNAME', '').strip()
        expected_password = os.getenv('WATCHDOG_AUTH_PASSWORD', '')
        if not expected_user or not expected_password:
            logger.error('REQUIRE_ACCESS_AUTH is enabled but WATCHDOG_AUTH_USERNAME/PASSWORD is not configured.')
            return Response('Access authentication is not configured.', status_code=503)

        def unauthorized():
            return Response(
                'Authentication required.', status_code=401,
                headers={'WWW-Authenticate': 'Basic realm="Fine-Print Watchdog", charset="UTF-8"'},
            )

        authorization = request.headers.get('authorization', '')
        try:
            scheme, encoded = authorization.split(' ', 1)
            if scheme.lower() != 'basic':
                return unauthorized()
            decoded = base64.b64decode(encoded, validate=True).decode('utf-8')
            username, password = decoded.split(':', 1)
        except (ValueError, UnicodeDecodeError, binascii.Error):
            return unauthorized()

        valid_user = hmac.compare_digest(username, expected_user)
        valid_password = hmac.compare_digest(password, expected_password)
        if not (valid_user and valid_password):
            return unauthorized()

    # A UUID cookie is a browser-scoped namespace, not an authentication token.
    # Authentication is checked separately above when enabled.
    raw_session = request.cookies.get('watchdog_session', '')
    try:
        session_id = str(uuid.UUID(raw_session))
        set_session_cookie = False
    except (ValueError, AttributeError):
        session_id = str(uuid.uuid4())
        set_session_cookie = not is_health
    request.state.session_id = session_id

    response = await call_next(request)
    if set_session_cookie:
        response.set_cookie(
            key='watchdog_session', value=session_id, max_age=8 * 60 * 60,
            httponly=True, secure=request.url.scheme == 'https', samesite='lax', path='/',
        )
    return response


def get_doc(key, session_id):
    with STORE_LOCK:
        item = DOCUMENTS.get(key)
    # Return the same response for missing and other-session documents to avoid
    # exposing whether a document ID belongs to another browser.
    if item is None or item.get('_owner_session') != session_id:
        raise HTTPException(404, 'Document not found in this browser session. Upload it again if needed.')
    return item


def public_document(item):
    """Never return the internal session owner marker to the browser."""
    return {key: value for key, value in item.items() if key != '_owner_session'}

def mode():
    value = os.getenv('WATCHDOG_MODE', 'demo')
    if value not in ('demo', 'local', 'hosted'):
        raise HTTPException(500, 'WATCHDOG_MODE must be demo, local or hosted.')
    return value

@app.get('/api/health')
def health():
    return {'status': 'ok', 'mode': mode(), 'model': llm.configured_model(), 'inference_checked': False}

@app.get('/api/rules')
def rules():
    return {'version': 'review-checklist-1', 'type': 'Practical review checklist; no legal authorities included', 'rules': RULES}

class PrivacyScanRequest(BaseModel):
    document_id: str

@app.post('/api/privacy/scan')
def privacy_scan(req: PrivacyScanRequest, request: Request):
    """Detect likely sensitive patterns without returning the matched values."""
    doc = get_doc(req.document_id, request.state.session_id)
    result = scan_pages(doc['pages'])
    result['document_id'] = doc['id']
    result['document_name'] = doc['name']
    return result

@app.get('/api/policies')
def policy_catalogue():
    """Return the curated India Policy Lens catalogue and official-source links."""
    return catalogue_payload()

class PolicyMatchRequest(BaseModel):
    document_id: str

@app.post('/api/policies/match')
def policy_lens_match(req: PolicyMatchRequest, request: Request):
    """Match uploaded text to catalogue keyword triggers inside its own browser session."""
    doc = get_doc(req.document_id, request.state.session_id)
    text = '\n'.join(page['text'] for page in doc['pages'])
    matches = match_policies(text)
    return {
        'document_id': doc['id'],
        'document_name': doc['name'],
        'document_sha256': doc['sha256'],
        'catalogue_version': CATALOGUE_VERSION,
        'catalogue_checked_on': CATALOGUE_CHECKED_ON,
        'catalogue_count': len(POLICIES),
        'matched_count': len(matches),
        'matches': matches,
        'matching_method': 'Literal keyword matching against extracted text. No AI or legal interpretation is used.',
        'warnings': [
            'A match is a research lead, not a conclusion that a law applies, that a clause is unlawful, or that the document complies.',
            'The catalogue is curated, not exhaustive. State/UT rules, Gazette notifications, amendments and sector-specific requirements may be missing.',
            'Check the official source for current amendments, commencement, scope, exceptions and the relevant transaction date.',
        ],
    }

@app.post('/api/documents')
async def upload(request: Request, file: UploadFile = File(...)):
    raw = await file.read(MAX_BYTES + 1)
    await file.close()
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, 'Maximum upload size is 8 MB.')
    ext = Path(file.filename or '').suffix.lower()
    warnings = []
    try:
        if ext in ('.txt', '.md'):
            pages = [{'page': 1, 'text': raw.decode('utf-8-sig')}]
        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError('Remove PDF password protection before uploading.')
            if len(reader.pages) > 80:
                raise ValueError('Maximum PDF length is 80 pages.')
            pages = [{'page': i+1, 'text': p.extract_text() or ''} for i,p in enumerate(reader.pages)]
            if any(not p['text'].strip() for p in pages):
                warnings.append('Some PDF pages have no extracted text. Images/scans are not reviewed; OCR is not included.')
        elif ext == '.docx':
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError('DOCX uncompressed contents exceed 32 MB.')
            document = Document(io.BytesIO(raw))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                lines.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
            pages = [{'page': 1, 'text': '\n'.join(lines)}]
            warnings.append('DOCX shown as one logical page; table text follows paragraphs. Headers, text boxes, tracked changes and images may be omitted.')
        else:
            raise ValueError('Supported formats: UTF-8 TXT/MD, text-based PDF, DOCX.')
    except Exception as exc:
        raise HTTPException(422, 'Unable to extract document: ' + str(exc)[:250])
    total = sum(len(p['text']) for p in pages)
    if total > MAX_CHARS:
        raise HTTPException(413, 'Maximum extracted length is 120,000 characters. Split the document.')
    if not any(p['text'].strip() for p in pages):
        raise HTTPException(422, 'No readable text. OCR scanned documents before uploading.')
    item = {'id': str(uuid.uuid4()), 'name': Path(file.filename or 'document').name, 'sha256': hashlib.sha256(raw).hexdigest(), 'pages': pages, 'characters': total, 'warnings': warnings, '_owner_session': request.state.session_id, '_created_at': time.time()}
    with STORE_LOCK:
        session_count = sum(1 for stored in DOCUMENTS.values() if stored.get('_owner_session') == request.state.session_id)
        if session_count >= MAX_DOCUMENTS_PER_SESSION:
            raise HTTPException(409, 'This browser session has reached its 20-document limit. Delete an uploaded document first.')
        if len(DOCUMENTS) >= MAX_DOCUMENTS_TOTAL:
            raise HTTPException(503, 'The temporary document store is at capacity. Please try again later.')
        DOCUMENTS[item['id']] = item
    return public_document(item)

@app.delete('/api/documents/{key}')
def delete(key: str, request: Request):
    get_doc(key, request.state.session_id)
    with STORE_LOCK:
        DOCUMENTS.pop(key, None)
    return {'deleted': True}

class AnalysisRequest(BaseModel):
    document_id: str
    language: Literal['English', 'German', 'Spanish', 'French', 'Italian', 'Japanese'] = 'English'
    privacy_categories: list[Literal['email', 'phone', 'pan', 'aadhaar', 'ifsc', 'bank_account']] = Field(default_factory=lambda: list(DEFAULT_PRIVACY_CATEGORIES), min_length=1)

@app.post('/api/analyze')
async def analyze(req: AnalysisRequest, request: Request):
    doc = get_doc(req.document_id, request.state.session_id)
    # Keep the original document for local evidence display, but only send the
    # same-length redacted copy to external inference. Stable offsets let the UI
    # highlight the corresponding original quotation without storing mappings.
    protected_pages, redaction_counts = redact_pages(doc['pages'], req.privacy_categories)
    findings = rule_findings(doc['pages'])
    rejected, failures = [], []
    parts = list(chunks(protected_pages))
    completed = 0
    if mode() != 'demo':
        for index, part in enumerate(parts):
            try:
                # Do not log contract text or model findings: uploaded documents may
                # contain confidential or personal information.
                proposed = await llm.extract(part, req.language)
                # Reject quotes lifted from other pages/chunks that were not supplied.
                accepted, refused = verify_candidates([x.model_dump() for x in proposed.findings], [part])
                for entry in accepted:
                    # Redacted text has the same character length, so offsets point
                    # to the corresponding range in the untouched source document.
                    entry['evidence'] = evidence(entry['quote'], protected_pages)
                findings.extend(accepted)
                rejected.extend(refused)
                completed += 1
            except Exception as exc:
                logger.exception('Nemotron analysis failed for chunk %s on page %s (%s)', index + 1, part['page'], type(exc).__name__)
                failures.append({'chunk': index+1, 'page': part['page'], 'error': (str(exc)[:220] if isinstance(exc, ValueError) else type(exc).__name__), 'action': 'Check provider/model availability and structured output; retry. The full exception is recorded in the server log without document text.'})
    unique = {}
    for finding in findings:
        # Chunk overlap can repeat the same finding. Normalize whitespace and case
        # before deduplication while retaining distinct titles for one clause.
        normalized_quote = ' '.join(str(finding.get('quote', '')).split()).casefold()
        normalized_title = ' '.join(str(finding.get('title', '')).split()).casefold()
        unique.setdefault((normalized_title, normalized_quote), finding)
    coverage_status = (
        'demo' if mode() == 'demo'
        else 'failed' if failures and completed == 0
        else 'partial' if failures
        else 'complete'
    )
    return {
        'document_id': doc['id'], 'document_sha256': doc['sha256'],
        'mode': mode(), 'language': req.language, 'findings': list(unique.values()),
        'rejected_findings': rejected, 'deadlines': deadline_candidates(doc['pages']),
        'privacy': {'enabled': True, 'categories': req.privacy_categories, 'masked_counts': redaction_counts,
                    'notice': 'Pattern-based masking protects only detected values in selected categories. Inspect the original document; this is not a guarantee that all personal data was found.'},
        'coverage': {
            'status': coverage_status, 'chunks_total': len(parts),
            'model_chunks_completed': completed, 'model_chunks_failed': len(failures),
            'failures': failures,
        },
        'warnings': doc['warnings'] + [
            'Exact source matching confirms quote presence only. Explanations and severity are unverified suggestions. No finding does not mean no risk.',
            'Checklist output is English; selected language applies to model explanations.',
            'Deadline detection covers explicit ISO dates on lines containing renewal/notice/expiry/cancellation terms only.',
        ],
    }

class CompareRequest(BaseModel):
    old_id: str
    new_id: str
    language: Literal['English', 'German', 'Spanish', 'French', 'Italian', 'Japanese'] = 'English'
    privacy_categories: list[Literal['email', 'phone', 'pan', 'aadhaar', 'ifsc', 'bank_account']] = Field(default_factory=lambda: list(DEFAULT_PRIVACY_CATEGORIES), min_length=1)

@app.post('/api/compare')
async def compare(req: CompareRequest, request: Request):
    a, b = get_doc(req.old_id, request.state.session_id), get_doc(req.new_id, request.state.session_id)
    ta, tb = '\n'.join(p['text'] for p in a['pages']), '\n'.join(p['text'] for p in b['pages'])
    protected_a, counts_a = redact_pages(a['pages'], req.privacy_categories)
    protected_b, counts_b = redact_pages(b['pages'], req.privacy_categories)
    protected_ta = '\n'.join(p['text'] for p in protected_a)
    protected_tb = '\n'.join(p['text'] for p in protected_b)
    changes = compare_text(ta, tb)  # local deterministic diff uses original text; never sent to the model
    candidates, rejected, warning = [], [], None
    if mode() != 'demo':
        if len(ta) + len(tb) > 10000:
            warning = 'Model contradiction pass skipped above 10,000 combined characters. Full deterministic text diff is included.'
        else:
            try:
                result = await llm.conflicts(protected_ta, protected_tb, req.language)
                for item in result.conflicts:
                    item = item.model_dump()
                    item['evidence_a'] = evidence(item['quote_a'], protected_a)
                    item['evidence_b'] = evidence(item['quote_b'], protected_b)
                    (candidates if item['evidence_a'] and item['evidence_b'] else rejected).append(item)
            except Exception as exc:
                logger.exception('Nemotron comparison failed (%s)', type(exc).__name__)
                reason = str(exc)[:180] if isinstance(exc, ValueError) else type(exc).__name__
                warning = 'Model comparison unavailable (' + reason + '). Text diff completed; inspect server logs for details.'
    return {'mode': mode(), 'old_sha256': a['sha256'], 'new_sha256': b['sha256'], 'changes': changes, 'possible_conflicts': candidates, 'rejected_conflicts': rejected, 'warning': warning, 'privacy': {'enabled': True, 'categories': req.privacy_categories, 'masked_counts_old': counts_a, 'masked_counts_new': counts_b, 'notice': 'Only selected pattern categories are masked before model comparison; inspect both source documents.'}, 'note': 'A text difference is not necessarily a contradiction or an unagreed change. Compare contract versions or brochure versus contract; review the full context.'}

class LoanRequest(BaseModel):
    principal: float = Field(gt=0, le=1e12, allow_inf_nan=False)
    annual_flat_percent: float = Field(ge=0, le=100, allow_inf_nan=False)
    months: int = Field(ge=1, le=600)
    upfront_fee: float = Field(default=0, ge=0, allow_inf_nan=False)
    @model_validator(mode='after')
    def check_fee(self):
        if self.upfront_fee >= self.principal * .95:
            raise ValueError('Fee must be less than 95% of principal.')
        return self

@app.post('/api/calculate')
def calculate(req: LoanRequest):
    return loan_cost(**req.model_dump())

class ContractCostRequest(BaseModel):
    contract_value: float = Field(default=0, ge=0, le=1e12, allow_inf_nan=False)
    recurring_payment: float = Field(default=0, ge=0, le=1e12, allow_inf_nan=False)
    recurring_periods: int = Field(default=0, ge=0, le=1200)
    upfront_charges: float = Field(default=0, ge=0, le=1e12, allow_inf_nan=False)
    penalties: float = Field(default=0, ge=0, le=1e12, allow_inf_nan=False)

@app.post('/api/calculate-contract')
def calculate_contract(req: ContractCostRequest):
    data = req.model_dump()
    recurring_total = data['recurring_payment'] * data['recurring_periods']
    estimated_total = data['contract_value'] + recurring_total + data['upfront_charges'] + data['penalties']
    return {
        'recurring_total': round(recurring_total, 2),
        'estimated_total': round(estimated_total, 2),
        'known_costs_total': round(data['upfront_charges'] + data['penalties'] + recurring_total, 2),
        'assumptions': 'Adds only the values entered. Recurring total is payment × number of periods. Taxes, escalations, refunds, conditional charges and unentered costs are excluded.'
    }

class DocumentRequest(BaseModel):
    document_id: str

@app.post('/api/deadlines')
def scan_deadlines(req: DocumentRequest, request: Request):
    doc = get_doc(req.document_id, request.state.session_id)
    return {'document_id': doc['id'], 'document_name': doc['name'], **document_deadline_candidates(doc['pages'])}

@app.post('/api/cost-candidates')
def scan_cost_candidates(req: DocumentRequest, request: Request):
    doc = get_doc(req.document_id, request.state.session_id)
    return {'document_id': doc['id'], 'document_name': doc['name'], **cost_candidates(doc['pages'])}

class CalendarRequest(BaseModel):
    title: str = Field(min_length=1, max_length=150)
    renewal_date: date
    notice_days: int = Field(ge=0, le=3650)
    description: str = Field(default='', max_length=2000)
    confirmed: bool

@app.post('/api/calendar')
def calendar_export(req: CalendarRequest):
    if not req.confirmed:
        raise HTTPException(422, 'Confirm the date and calendar-day notice period before export.')
    try:
        text = make_calendar(req.title, req.renewal_date, req.notice_days, req.description)
    except OverflowError:
        raise HTTPException(422, 'Date is outside the supported calendar range.')
    return Response(text, media_type='text/calendar', headers={'Content-Disposition': 'attachment; filename="notice-deadline.ics"'})

DIST = Path(__file__).resolve().parents[2] / 'frontend' / 'dist'
if DIST.exists():
    app.mount('/', StaticFiles(directory=DIST, html=True), name='frontend')
