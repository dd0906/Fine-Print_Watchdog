import asyncio
import base64
import io
import json
from datetime import date
import httpx
import pytest
from fastapi.testclient import TestClient
from app import main, llm
from app.core import chunks, evidence, loan_cost, make_calendar, verify_candidates

@pytest.fixture(autouse=True)
def clean(monkeypatch):
    main.DOCUMENTS.clear()
    monkeypatch.setenv('WATCHDOG_MODE', 'demo')
    monkeypatch.setenv('REQUIRE_ACCESS_AUTH', 'false')

@pytest.fixture
def client():
    return TestClient(main.app)

def upload(client, text, name='test.txt'):
    response = client.post('/api/documents', files={'file': (name,text.encode(),'text/plain')})
    assert response.status_code == 200, response.text
    return response.json()

def test_exact_evidence_rejects_fabrication():
    pages=[{'page':2,'text':'The early termination fee is 100. The early termination fee is 100.'}]
    assert len(evidence('The early termination fee is 100.',pages))==2
    assert not evidence('The early termination fee is 200.',pages)
    assert not evidence('fee',pages)

def test_loan_independent_cash_flow():
    r=loan_cost(100000,12,12,2000)
    assert r['monthly_payment']==9333.33
    assert r['total_cost_including_fee']==14000
    monthly=(1+r['effective_annual_percent']/100)**(1/12)-1
    assert sum((112000/12)/(1+monthly)**i for i in range(1,13))==pytest.approx(98000,abs=.1)
    assert loan_cost(1200,0,12)['effective_annual_percent']==0

@pytest.mark.parametrize('body',[
    {'principal':100,'annual_flat_percent':12,'months':0},
    {'principal':100,'annual_flat_percent':12,'months':12,'upfront_fee':100},
    {'principal':-1,'annual_flat_percent':12,'months':12},
])
def test_invalid_finance(client,body):
    assert client.post('/api/calculate',json=body).status_code==422

def test_calendar_math_and_injection():
    text=make_calendar('Notice\nBEGIN:VEVENT',date(2028,3,1),1,'a,b;c')
    assert 'DTSTART;VALUE=DATE:20280229' in text
    assert 'SUMMARY:Notice\\nBEGIN:VEVENT' in text
    assert 'DESCRIPTION:a\\,b\\;c' in text
    assert text.count('\r\nBEGIN:VEVENT\r\n')==1

def test_calendar_confirmation(client):
    body={'title':'Notice','renewal_date':'2027-01-31','notice_days':30,'confirmed':False}
    assert client.post('/api/calendar',json=body).status_code==422
    body['confirmed']=True
    r=client.post('/api/calendar',json=body)
    assert r.status_code==200 and 'DTSTART;VALUE=DATE:20270101' in r.text

def test_policy_catalogue_has_official_sources(client):
    response = client.get('/api/policies')
    assert response.status_code == 200
    data = response.json()
    assert data['count'] >= 40
    assert 'Finance & lending' in data['categories']
    assert all(item['source_url'].startswith('https://') for item in data['items'])
    assert any('Digital Personal Data Protection Act' in item['title'] for item in data['items'])
    assert 'not an exhaustive database' in data['notice']


def test_policy_lens_matches_loan_terms(client):
    doc = upload(client, 'Digital loan agreement. Principal 100000, processing fee 2000, annual percentage rate APR, EMI repayment and lender disclosure.')
    response = client.post('/api/policies/match', json={'document_id': doc['id']})
    assert response.status_code == 200
    result = response.json()
    titles = [item['title'] for item in result['matches']]
    assert 'Reserve Bank of India (Digital Lending) Directions, 2025' in titles
    assert 'RBI — Key Facts Statement (KFS) for Loans & Advances (15 April 2024)' in titles
    assert result['matched_count'] == len(result['matches'])
    assert any('keyword' in warning.lower() or 'match is a research lead' in warning.lower() for warning in result['warnings'])


def test_policy_lens_respects_browser_document_isolation(monkeypatch):
    monkeypatch.setenv('REQUIRE_ACCESS_AUTH', 'true')
    monkeypatch.setenv('WATCHDOG_AUTH_USERNAME', 'watchdog')
    monkeypatch.setenv('WATCHDOG_AUTH_PASSWORD', 'test-password-strong')
    headers = {'Authorization': 'Basic ' + base64.b64encode(b'watchdog:test-password-strong').decode('ascii')}
    one, two = TestClient(main.app), TestClient(main.app)
    uploaded = one.post('/api/documents', files={'file': ('lease.txt', b'Lease agreement with a rental deposit and notice period.', 'text/plain')}, headers=headers)
    assert uploaded.status_code == 200, uploaded.text
    doc = uploaded.json()
    assert two.post('/api/policies/match', json={'document_id': doc['id']}, headers=headers).status_code == 404
    one.close(); two.close()


def test_demo_end_to_end(client):
    doc=upload(client,'The agreement automatically renews on 2027-01-31.\nEarly termination costs 500.')
    report=client.post('/api/analyze',json={'document_id':doc['id']}).json()
    assert report['mode']=='demo' and len(report['findings'])==2
    assert report['coverage']['model_chunks_completed']==0
    assert report['deadlines'][0]['date']=='2027-01-31'
    assert all(f['evidence'] for f in report['findings'])
    other=upload(client,'The agreement automatically renews on 2027-01-31.\nEarly termination costs 900.')
    comparison=client.post('/api/compare',json={'old_id':doc['id'],'new_id':other['id']}).json()
    assert comparison['changes'][0]['after']=='Early termination costs 900.'
    assert client.delete('/api/documents/'+doc['id']).status_code==200
    assert client.post('/api/analyze',json={'document_id':doc['id']}).status_code==404

def test_model_rejections_and_failure_visibility(client,monkeypatch):
    monkeypatch.setenv('WATCHDOG_MODE','local')
    doc=upload(client,'The tenant must pay rent monthly.')
    async def fake(*args):
        return llm.Findings(findings=[llm.Finding(title='Rent',quote='The tenant must pay rent monthly.',explanation='Review timing',question='Which day?'),llm.Finding(title='Fake',quote='The tenant must pay rent weekly.',explanation='Invented',question='Which day?')])
    monkeypatch.setattr(llm,'extract',fake)
    result=client.post('/api/analyze',json={'document_id':doc['id']}).json()
    assert len(result['findings'])==1 and len(result['rejected_findings'])==1
    assert result['coverage']['model_chunks_completed']==1
    assert result['coverage']['status']=='complete'
    async def fail(*args):
        raise ValueError('bad JSON')
    monkeypatch.setattr(llm,'extract',fail)
    result=client.post('/api/analyze',json={'document_id':doc['id']}).json()
    assert len(result['coverage']['failures'])==1
    assert result['coverage']['model_chunks_completed']==0
    assert result['coverage']['status']=='failed'
    assert result['coverage']['model_chunks_failed']==1

def test_model_conflicts_verify_both_sources(client,monkeypatch):
    monkeypatch.setenv('WATCHDOG_MODE','local')
    a=upload(client,'Cancellation is free at any time.')
    b=upload(client,'Cancellation requires a fee of 500.')
    async def fake(*args):
        return llm.Conflicts(conflicts=[llm.Conflict(quote_a='Cancellation is free at any time.',quote_b='Cancellation requires a fee of 500.',explanation='Possible discrepancy'),llm.Conflict(quote_a='Cancellation is free at any time.',quote_b='Cancellation requires a fee of 900.',explanation='Fabricated')])
    monkeypatch.setattr(llm,'conflicts',fake)
    r=client.post('/api/compare',json={'old_id':a['id'],'new_id':b['id']}).json()
    assert len(r['possible_conflicts'])==1 and len(r['rejected_conflicts'])==1

def test_upload_validation(client):
    assert client.post('/api/documents',files={'file':('a.exe',b'data')}).status_code==422
    assert client.post('/api/documents',files={'file':('a.txt',b'   ')}).status_code==422
    assert client.post('/api/documents',files={'file':('a.txt',b'a'*120001)}).status_code==413

def test_docx_table_extraction(client):
    from docx import Document
    d=Document();d.add_paragraph('A sample contract.')
    t=d.add_table(rows=1,cols=1);t.cell(0,0).text='A non-refundable processing fee applies.'
    b=io.BytesIO();d.save(b)
    r=client.post('/api/documents',files={'file':('contract.docx',b.getvalue())})
    assert r.status_code==200 and 'processing fee' in r.json()['pages'][0]['text']

def test_chunk_coverage():
    text=''.join(str(i%10) for i in range(21000))
    parts=list(chunks([{'page':1,'text':text}]))
    covered=set()
    for p in parts:
        covered.update(range(p['start'],p['start']+len(p['text'])))
    assert covered==set(range(len(text)))

def test_protocol_parsing(monkeypatch):
    real_client = httpx.AsyncClient

    def handler(request):
        payload = json.loads(request.content)
        assert payload['messages'][0]['content'].startswith('/no_think')
        assert payload['stream'] is True

        if llm._provider_config()['provider'] == 'nebius':
            assert 'chat_template_kwargs' not in payload
        else:
            assert payload['chat_template_kwargs']['enable_thinking'] is False

        events = [
            {
                'choices': [
                    {
                        'index': 0,
                        'delta': {
                            'content': '```json\n{"findings": []}\n```'
                        },
                        'finish_reason': None,
                    }
                ]
            },
            {
                'choices': [
                    {
                        'index': 0,
                        'delta': {},
                        'finish_reason': 'stop',
                    }
                ]
            },
        ]

        body = ''.join(
            'data: ' + json.dumps(event) + '\n\n'
            for event in events
        ) + 'data: [DONE]\n\n'

        return httpx.Response(
            200,
            headers={'content-type': 'text/event-stream'},
            content=body.encode(),
        )

    transport = httpx.MockTransport(handler)

    monkeypatch.setattr(
        llm.httpx,
        'AsyncClient',
        lambda **kw: real_client(transport=transport, **kw),
    )

    assert asyncio.run(
        llm.request_json('data', llm.Findings)
    ).findings == []

def test_basic_auth_guard(client, monkeypatch):
    monkeypatch.setenv('REQUIRE_ACCESS_AUTH', 'true')
    monkeypatch.setenv('WATCHDOG_AUTH_USERNAME', 'watchdog')
    monkeypatch.setenv('WATCHDOG_AUTH_PASSWORD', 'test-password-strong')

    assert client.get('/api/rules').status_code == 401
    # Render's platform health check must keep working without browser credentials.
    assert client.get('/api/health').status_code == 200

    token = base64.b64encode(b'watchdog:test-password-strong').decode('ascii')
    response = client.get('/api/rules', headers={'Authorization': 'Basic ' + token})
    assert response.status_code == 200

    wrong = base64.b64encode(b'watchdog:wrong').decode('ascii')
    assert client.get('/api/rules', headers={'Authorization': 'Basic ' + wrong}).status_code == 401


def test_basic_auth_fails_closed_when_password_missing(client, monkeypatch):
    monkeypatch.setenv('REQUIRE_ACCESS_AUTH', 'true')
    monkeypatch.setenv('WATCHDOG_AUTH_USERNAME', 'watchdog')
    monkeypatch.delenv('WATCHDOG_AUTH_PASSWORD', raising=False)
    assert client.get('/api/rules').status_code == 503

def test_pdf_blank_rejected(client):
    from pypdf import PdfWriter
    writer=PdfWriter();writer.add_blank_page(width=100,height=100)
    data=io.BytesIO();writer.write(data)
    assert client.post('/api/documents',files={'file':('blank.pdf',data.getvalue())}).status_code==422


def test_all_sensitive_routes_require_auth(client, monkeypatch):
    monkeypatch.setenv('REQUIRE_ACCESS_AUTH', 'true')
    monkeypatch.setenv('WATCHDOG_AUTH_USERNAME', 'watchdog')
    monkeypatch.setenv('WATCHDOG_AUTH_PASSWORD', 'test-password-strong')
    unauthenticated_calls = [
        lambda: client.get('/'),
        lambda: client.get('/api/rules'),
        lambda: client.get('/api/policies'),
        lambda: client.post('/api/policies/match', json={'document_id': 'not-a-real-id'}),
        lambda: client.post('/api/documents', files={'file': ('a.txt', b'private contract text') }),
        lambda: client.delete('/api/documents/not-a-real-id'),
        lambda: client.post('/api/analyze', json={'document_id': 'not-a-real-id'}),
        lambda: client.post('/api/compare', json={'old_id': 'a', 'new_id': 'b'}),
        lambda: client.post('/api/calculate', json={'principal': 100, 'annual_flat_percent': 5, 'months': 12}),
        lambda: client.post('/api/calendar', json={'title': 'Notice', 'renewal_date': '2027-01-31', 'notice_days': 30, 'confirmed': True}),
    ]
    for call in unauthenticated_calls:
        response = call()
        assert response.status_code == 401, (response.request.method, response.request.url, response.status_code)


def test_documents_are_isolated_between_browser_sessions(monkeypatch):
    monkeypatch.setenv('REQUIRE_ACCESS_AUTH', 'true')
    monkeypatch.setenv('WATCHDOG_AUTH_USERNAME', 'watchdog')
    monkeypatch.setenv('WATCHDOG_AUTH_PASSWORD', 'test-password-strong')
    token = base64.b64encode(b'watchdog:test-password-strong').decode('ascii')
    headers = {'Authorization': 'Basic ' + token}
    browser_a = TestClient(main.app)
    browser_b = TestClient(main.app)

    uploaded = browser_a.post('/api/documents', headers=headers, files={'file': ('private.txt', b'The tenant may terminate this agreement with written notice.')})
    assert uploaded.status_code == 200, uploaded.text
    doc_id = uploaded.json()['id']
    assert ' _owner_session' not in uploaded.text
    assert '_owner_session' not in uploaded.text

    # The second browser uses the same shared Basic Auth credentials, but its
    # HttpOnly session cookie gives it a separate document namespace.
    denied_review = browser_b.post('/api/analyze', headers=headers, json={'document_id': doc_id})
    assert denied_review.status_code == 404
    denied_delete = browser_b.delete('/api/documents/' + doc_id, headers=headers)
    assert denied_delete.status_code == 404
    assert browser_a.post('/api/analyze', headers=headers, json={'document_id': doc_id}).status_code == 200


def test_analysis_coverage_reports_failed_and_partial(monkeypatch, client):
    monkeypatch.setenv('WATCHDOG_MODE', 'hosted')
    doc = upload(client, 'The tenant must pay rent monthly.')
    calls = {'count': 0}

    async def partly_fail(*args):
        calls['count'] += 1
        if calls['count'] == 1:
            return llm.Findings(findings=[])
        raise TimeoutError('synthetic timeout')

    monkeypatch.setattr(llm, 'extract', partly_fail)
    # A document longer than one chunk produces more than one model passage.
    long_text = ('The tenant must pay rent monthly. ' + ('Ordinary terms apply. ' * 500))
    main.DOCUMENTS[doc['id']]['pages'] = [{'page': 1, 'text': long_text}]
    result = client.post('/api/analyze', json={'document_id': doc['id']}).json()
    assert result['coverage']['chunks_total'] > 1
    assert result['coverage']['status'] == 'partial'
    assert result['coverage']['model_chunks_completed'] >= 1
    assert result['coverage']['model_chunks_failed'] >= 1


def test_overlap_duplicates_are_deduplicated(client, monkeypatch):
    monkeypatch.setenv('WATCHDOG_MODE', 'local')
    quote = 'The supplier may increase fees after giving written notice.'
    text = ('Ordinary wording without special terms. ' * 150) + quote + (' Routine text continues. ' * 200)
    doc = upload(client, text)

    async def same_finding(*args):
        return llm.Findings(findings=[llm.Finding(title='Fee increase', quote=quote, explanation='Check the notice period.', question='How much notice is required?')])

    monkeypatch.setattr(llm, 'extract', same_finding)
    result = client.post('/api/analyze', json={'document_id': doc['id']}).json()
    model_findings = [x for x in result['findings'] if x.get('origin') == 'model suggestion']
    assert result['coverage']['chunks_total'] > 1
    assert len(model_findings) == 1


def test_transient_model_transport_error_retries_once(monkeypatch):
    real_client = httpx.AsyncClient
    state = {'calls': 0}

    def handler(request):
        state['calls'] += 1
        if state['calls'] == 1:
            raise httpx.ReadTimeout('temporary timeout', request=request)
        event = {'choices': [{'index': 0, 'delta': {'content': '{"findings": []}'}, 'finish_reason': 'stop'}]}
        body = 'data: ' + json.dumps(event) + '\n\n' + 'data: [DONE]\n\n'
        return httpx.Response(200, headers={'content-type': 'text/event-stream'}, content=body.encode())

    transport = httpx.MockTransport(handler)
    monkeypatch.setattr(llm.httpx, 'AsyncClient', lambda **kw: real_client(transport=transport, **kw))
    result = asyncio.run(llm.request_json('retry test', llm.Findings))
    assert result.findings == []
    assert state['calls'] == 2


def test_documents_expire_from_temporary_store(client):
    doc = upload(client, 'The agreement automatically renews on 2027-01-31.')
    main.DOCUMENTS[doc['id']]['_created_at'] = 0
    # Every request, including health checks, triggers expiry cleanup.
    assert client.get('/api/health').status_code == 200
    assert doc['id'] not in main.DOCUMENTS
    assert client.post('/api/analyze', json={'document_id': doc['id']}).status_code == 404


def test_demo_mode_never_calls_model_even_if_nebius_key_is_present(client, monkeypatch):
    monkeypatch.setenv('WATCHDOG_MODE', 'demo')
    monkeypatch.setenv('NEBIUS_API_KEY', 'test-key-must-not-be-used')

    async def forbidden(*args, **kwargs):
        raise AssertionError('Demo mode must never call the model provider')

    monkeypatch.setattr(main.llm, 'extract', forbidden)
    doc = upload(client, 'The agreement automatically renews on 2027-01-31. Early termination costs 500.')
    response = client.post('/api/analyze', json={'document_id': doc['id']})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['mode'] == 'demo'
    assert result['coverage']['status'] == 'demo'
    assert result['coverage']['model_chunks_completed'] == 0


def test_deadline_scan_is_independent_and_finds_notice_period(client):
    doc = upload(client, 'Either party may terminate by giving 45 calendar days notice. The agreement renews on 2027-06-30.')
    response = client.post('/api/deadlines', json={'document_id': doc['id']})
    assert response.status_code == 200
    data = response.json()
    assert data['dates'][0]['date'] == '2027-06-30'
    assert data['notice_periods'][0]['stated_value'] == 45
    assert data['notice_periods'][0]['unit'] == 'calendar days'
    assert '45 calendar days notice' in data['notice_periods'][0]['quote']


def test_cost_candidates_detect_loan_and_contract_figures(client):
    text = ('Loan amount: Rs. 150,000. Annual flat interest rate: 12%. '
            'Loan tenure: 24 months. Processing fee: Rs. 1,500. '
            'Contract value: Rs. 200,000. Monthly rent: Rs. 12,000.')
    doc = upload(client, text)
    response = client.post('/api/cost-candidates', json={'document_id': doc['id']})
    assert response.status_code == 200
    data = response.json()
    assert data['kind'] == 'mixed'
    values = {(item['field'], item['value']) for item in data['candidates']}
    assert ('principal', 150000) in values
    assert ('annual_flat_percent', 12) in values
    assert ('months', 24) in values
    assert ('upfront_fee', 1500) in values
    assert ('contract_value', 200000) in values
    assert ('recurring_payment', 12000) in values


def test_general_contract_cost_calculation(client):
    response = client.post('/api/calculate-contract', json={
        'contract_value': 100000, 'recurring_payment': 2000,
        'recurring_periods': 12, 'upfront_charges': 500, 'penalties': 0,
    })
    assert response.status_code == 200
    data = response.json()
    assert data['recurring_total'] == 24000
    assert data['estimated_total'] == 124500
    assert data['known_costs_total'] == 24500
